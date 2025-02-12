from django.core.management.base import BaseCommand
from myapp.models import (
    Farm,
    SubFarm,
    Deforestation,
    ERDA,
)
import openpyxl

class Command(BaseCommand):
    def handle(self, *args, **options):
        print('hola mundo desde django')
        # Cargar el archivo Excel
        archivo = "C:/Users/Usuario/Dev/__farm/abrir_archivo\Programa_python.xlsx"  # Reemplaza con el nombre de tu archivo
        wb = openpyxl.load_workbook(archivo)
        
        hoja = 'Farm_data'
        ws = wb[hoja]  # Obtener la hoja
        primera_fila = True
        for fila in ws.iter_rows(values_only=True):  # `values_only=True` devuelve solo los valores
            try: 
                if primera_fila:
                    primera_fila = False
                else:
                    num_int = fila[0]  or "N/A"
                    farm_id = fila[1] or "N/A"
                    tecnico = fila[2] or "N/A"
                    map_date = fila[3] or "N/A"
                    region = fila[4] or "N/A"
                    nom_pol = fila[5] or "N/A"
                    area_pol = fila[6] or 0
                    latitud = fila[7] or "N/A"
                    longitud = fila[8] or "N/A"

                    fd = Farm_data(
                        num_int=num_int,
                        farm_id=farm_id,
                        tecnico=tecnico,
                        map_date=map_date,
                        region=region,
                        nom_pol=nom_pol,
                        area_pol=area_pol,
                        latitud=latitud,
                        longitud=longitud,
                    )

                fd.save()
            except:
                print('fallo: {} fila: {}'.format(hoja, fila))
        hoja = 'Datos_totales'
        ws = wb[hoja]  # Obtener la hoja
        primera_fila = True
        for fila in ws.iter_rows(values_only=True):  # `values_only=True` devuelve solo los valores
            try: 
                if primera_fila:
                    primera_fila = False
                else:
                    num_int = fila[0]  or "N/A"
                    region = fila[1] or "N/A"
                    propietario = fila[2] or "N/A"
                    municipio = fila[3] or "N/A"
                    departamento = fila[4] or "N/A"

                    fd = Datos_totales(
                        num_int =num_int, 
                        region =region, 
                        propietario =propietario, 
                        municipio =municipio, 
                        departamento =departamento, 
                    )

                    fd.save()

                    num_int = fila[0]  or "N/A"
                    farm_id = fila[1] or "N/A"
                    tecnico = fila[2] or "N/A"
                    map_date = fila[3] or "N/A"
                    region = fila[4] or "N/A"
                    nom_pol = fila[5] or "N/A"
                    area_pol = fila[6] or 0
                    latitud = fila[7] or "N/A"
                    longitud = fila[8] or "N/A"

                    fd = Farm_data(
                        num_int=num_int,
                        farm_id=farm_id,
                        tecnico=tecnico,
                        map_date=map_date,
                        region=region,
                        nom_pol=nom_pol,
                        area_pol=area_pol,
                        latitud=latitud,
                        longitud=longitud,
                    )

                    fd.save()
            except:
                print('fallo: {} fila: {}'.format(hoja, fila))
        hoja = 'Deforestación'
        ws = wb[hoja]  # Obtener la hoja
        primera_fila = True
        for fila in ws.iter_rows(values_only=True):  # `values_only=True` devuelve solo los valores
            try:
                if primera_fila:
                    primera_fila = False
                else:
                    num_int = fila[0]  or "N/A"
                    farm_id = fila[2] or "N/A"
                    nivel = fila[3] or "N/A"
                    Fase = fila[4] or "N/A"

                    fd = Deforestacion(
                        num_int=num_int,
                        farm_id=farm_id,
                        nivel=nivel,
                        Fase=Fase,
                    )

                    fd.save()               
                    num_int = fila[0]  or "N/A"
                    farm_id = fila[1] or "N/A"
                    tecnico = fila[2] or "N/A"
                    map_date = fila[3] or "N/A"
                    region = fila[4] or "N/A"
                    nom_pol = fila[5] or "N/A"
                    area_pol = fila[6] or 0
                    latitud = fila[7] or "N/A"
                    longitud = fila[8] or "N/A"

                    fd = Farm_data(
                        num_int=num_int,
                        farm_id=farm_id,
                        tecnico=tecnico,
                        map_date=map_date,
                        region=region,
                        nom_pol=nom_pol,
                        area_pol=area_pol,
                        latitud=latitud,
                        longitud=longitud,
                    )

                    fd.save()
            except:
                print('fallo: {} fila: {}'.format(hoja, fila))
        hoja = 'ERDA'
        ws = wb[hoja]  # Obtener la hoja
        primera_fila = True
        for fila in ws.iter_rows(values_only=True):  # `values_only=True` devuelve solo los valores
            try:
                if primera_fila:
                    primera_fila = False
                else:
                    num_int = fila[6]  or "N/A"
                    farm_id = fila[8] or "N/A"
                    risk = fila[23] or "N/A"

                    fd = ERDA(
                        num_int=num_int,
                        farm_id=farm_id,
                        risk=risk,
                    )

                    fd.save()
            except:
                print('fallo: {} fila: {}'.format(hoja, fila))


        # Cerrar el archivo
        wb.close()